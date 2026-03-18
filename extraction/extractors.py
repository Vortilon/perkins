from __future__ import annotations

from io import BytesIO
from typing import Iterable


def _normalize(text: str) -> str:
    # Keep it simple: collapse trailing spaces and normalize newlines.
    return "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")).strip()


def extract_text(filename: str, content_type: str | None, data: bytes) -> str:
    """
    Extract text from common customer report formats.

    Supported:
    - PDF (text-based): .pdf / application/pdf
    - Word: .docx / application/vnd.openxmlformats-officedocument.wordprocessingml.document
    - Excel: .xlsx / application/vnd.openxmlformats-officedocument.spreadsheetml.sheet
    - Plain text: .txt / text/plain

    Note: scanned PDFs (images) will return little/no text without OCR.
    """
    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    if name.endswith(".txt") or ctype.startswith("text/plain"):
        return _normalize(data.decode("utf-8", errors="replace"))

    if name.endswith(".pdf") or ctype == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        parts: list[str] = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return _normalize("\n\n".join(parts))

    if name.endswith(".docx") or ctype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        from docx import Document

        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        return _normalize("\n".join(parts))

    if name.endswith(".xlsx") or ctype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
        chunks: list[str] = []
        for ws in wb.worksheets:
            chunks.append(f"## Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if not row:
                    continue
                vals = ["" if v is None else str(v) for v in row]
                if any(v.strip() for v in vals):
                    chunks.append("\t".join(vals))
            chunks.append("")
        return _normalize("\n".join(chunks))

    raise ValueError(f"Unsupported file type: filename={filename!r}, content_type={content_type!r}")

